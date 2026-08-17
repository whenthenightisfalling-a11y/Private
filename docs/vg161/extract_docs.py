"""把 .docx / .pptx 提取为纯文本，便于审阅与检索。

用途：VG161 项目的方案、SAP skeleton、审评回复与 CDE 沟通材料多为 Word/PPT
格式，且 SAP skeleton 的关键信息（终点定义、α 分配、样本量假设、中期分析）
大量位于表格中。本脚本按顺序提取段落与表格，避免遗漏表格内容。

依赖：
    pip install python-docx python-pptx

用法：
    python3 docs/vg161/extract_docs.py <输入目录> [-o <输出目录>]

输出：每个源文件对应一个同名 .txt，默认写入 <输入目录>/_extracted/。
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path


def extract_docx(path: Path) -> str:
    """提取 Word 文档的段落与表格，保持文档内的原始顺序。"""
    import docx
    from docx.document import Document
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = docx.Document(str(path))
    lines: list[str] = []

    def walk(parent: Document):
        """按 XML 顺序遍历 body，使段落与表格不错位。"""
        for child in parent.element.body.iterchildren():
            if isinstance(child, CT_P):
                yield Paragraph(child, parent)
            elif isinstance(child, CT_Tbl):
                yield Table(child, parent)

    for block in walk(document):
        if isinstance(block, Paragraph):
            text = block.text.strip()
            if not text:
                continue
            style = (block.style.name or "").lower()
            if style.startswith("heading"):
                level = "".join(c for c in style if c.isdigit()) or "1"
                lines.append(f"\n{'#' * min(int(level), 6)} {text}\n")
            else:
                lines.append(text)
        else:
            lines.append("\n[表格]")
            for row in block.rows:
                cells = [c.text.strip().replace("\n", " ") for c in row.cells]
                lines.append("| " + " | ".join(cells) + " |")
            lines.append("")

    return "\n".join(lines)


def extract_pptx(path: Path) -> str:
    """提取 PPT 每页的文本框、表格与备注。"""
    from pptx import Presentation

    presentation = Presentation(str(path))
    lines: list[str] = []

    for index, slide in enumerate(presentation.slides, start=1):
        lines.append(f"\n===== 第 {index} 页 =====")
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    lines.append(text)
            if getattr(shape, "has_table", False):
                lines.append("[表格]")
                for row in shape.table.rows:
                    cells = [c.text.strip().replace("\n", " ") for c in row.cells]
                    lines.append("| " + " | ".join(cells) + " |")
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append(f"[备注] {notes}")

    return "\n".join(lines)


EXTRACTORS = {".docx": extract_docx, ".pptx": extract_pptx}


def main() -> int:
    parser = argparse.ArgumentParser(description="批量提取 docx/pptx 文本")
    parser.add_argument("source", type=Path, help="包含 docx/pptx 的目录")
    parser.add_argument("-o", "--out", type=Path, help="输出目录")
    args = parser.parse_args()

    if not args.source.is_dir():
        print(f"错误：{args.source} 不是目录", file=sys.stderr)
        return 1

    out_dir = args.out or args.source / "_extracted"
    out_dir.mkdir(parents=True, exist_ok=True)

    files = sorted(
        p
        for p in args.source.rglob("*")
        if p.suffix.lower() in EXTRACTORS and not p.name.startswith("~$")
    )
    if not files:
        print(f"在 {args.source} 下未找到 docx/pptx 文件")
        return 1

    for path in files:
        try:
            text = EXTRACTORS[path.suffix.lower()](path)
        except Exception as exc:  # 单个文件损坏不应中断整批提取
            print(f"跳过 {path.name}：{exc}", file=sys.stderr)
            continue
        target = out_dir / f"{path.stem}.txt"
        target.write_text(text, encoding="utf-8")
        print(f"{path.name} -> {target.name}（{len(text.splitlines())} 行）")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
